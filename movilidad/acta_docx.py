"""Genera el docx del acta «After the Mobility» (RF-17), adaptado de lib_docx.py (skill
acta-reconocimiento): relleno de la plantilla con python-docx + cirugía de XML para la cabecera
(nombre y curso) y el bloque de firma.

Los recursos (plantilla_simplificado.docx, header1_template.xml, erasmus_logo.png) no están en el
repositorio: se leen de settings.RECURSOS_ACTA_DIR (la carpeta ToRs26/_recursos_acta del
coordinador). El PDF solo se genera si LibreOffice (`soffice`) está instalado.
"""

import io
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

import docx
from django.conf import settings
from django.core.files.base import ContentFile
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

FIRMA_XML = (
    '<w:p><w:pPr><w:rPr><w:lang w:val="es-ES"/></w:rPr></w:pPr></w:p>'
    '<w:p><w:pPr><w:jc w:val="right"/><w:rPr><w:u w:val="single"/></w:rPr></w:pPr>'
    '<w:r><w:t xml:space="preserve">A Coruña, en la fecha que figura en la firma digital</w:t></w:r></w:p>'
    '<w:p><w:pPr><w:jc w:val="right"/><w:rPr><w:lang w:val="es-ES"/></w:rPr></w:pPr>'
    '<w:r><w:rPr><w:lang w:val="es-ES"/></w:rPr><w:t xml:space="preserve">Asdo: {firmante}</w:t></w:r></w:p>'
    '<w:p><w:pPr><w:jc w:val="right"/><w:rPr><w:lang w:val="es-ES"/></w:rPr></w:pPr>'
    '<w:r><w:rPr><w:lang w:val="es-ES"/></w:rPr><w:t>{cargo}</w:t></w:r></w:p>'
)

# Patrón habitual de la plantilla «Simplificado»: dos párrafos vacíos tras la tabla.
COLA_PLANTILLA = re.compile(
    r'<w:p [^>]*w:rsidR="00D363A9"[^>]*>.*?</w:p><w:p [^>]*w:rsidR="000726FD"[^>]*w:rsidP="003D37D2"[^>]*>.*?</w:p>'
    r'(?=<w:sectPr)', re.DOTALL)


class RecursosNoDisponibles(RuntimeError):
    pass


def recurso(nombre) -> Path:
    base = Path(settings.RECURSOS_ACTA_DIR)
    ruta = base / nombre
    if not ruta.exists():
        raise RecursosNoDisponibles(
            f"No encuentro {nombre} en {base}. Configura RECURSOS_ACTA_DIR con la carpeta "
            f"ToRs26/_recursos_acta (plantilla_simplificado.docx, header1_template.xml, erasmus_logo.png)."
        )
    return ruta


def curso_largo(curso: str) -> str:
    """'2026/27' -> '2026/2027'."""
    m = re.fullmatch(r"(\d{4})/(\d{2,4})", curso or "")
    if not m:
        return curso
    fin = m.group(2) if len(m.group(2)) == 4 else m.group(1)[:2] + m.group(2)
    return f"{m.group(1)}/{fin}"


def reemplazar_texto_partido(xml: str, patron: str, nuevo: str, despues_de: str = "") -> str:
    """Sustituye en el texto visible de un XML (aunque Word lo haya partido en varios <w:t>) cada
    coincidencia de `patron` que siga a una aparición de `despues_de`, carácter a carácter (las
    plantillas repiten la cabecera en su versión alternativa). `nuevo` debe tener la misma longitud."""
    posiciones, texto, dentro = [], [], False
    for i, c in enumerate(xml):
        if c == "<":
            dentro = True
        elif c == ">":
            dentro = False
        elif not dentro:
            posiciones.append(i)
            texto.append(c)
    texto = "".join(texto)
    inicios = [m.end() for m in re.finditer(re.escape(despues_de), texto)] if despues_de else [0]
    chars = list(xml)
    for inicio in inicios:
        m = re.compile(patron).search(texto, inicio)
        if not m or len(m.group(0)) != len(nuevo):
            continue
        for k, c in zip(range(m.start(), m.end()), nuevo):
            chars[posiciones[k]] = c
    return "".join(chars)


def _celda(fila, col, valor, negrita=False, alinear=WD_ALIGN_PARAGRAPH.CENTER):
    celda = fila.cells[col]
    celda.text = str(valor)
    for p in celda.paragraphs:
        p.alignment = alinear
        pf = p.paragraph_format
        pf.space_after = pf.space_before = Pt(0)
        pf.line_spacing = 1.0
        for r in p.runs:
            r.font.size, r.font.name, r.bold = Pt(8), "Calibri", negrita


def localizar_filas(documento) -> dict:
    idx = {}
    for i, fila in enumerate(documento.tables[0].rows):
        textos = [c.text.strip() for c in fila.cells]
        junto = " ".join(textos)
        if "Component code" in junto and "Receiving Institution" in junto and "Grades received" in junto:
            idx["cab_c"] = i
        elif "Component code" in junto and "recognised" in junto and "Grades registered" in junto:
            idx["cab_d"] = i
        elif textos[0].startswith("Table C") and any(t.startswith("Total") for t in textos[1:]):
            idx["total_c"] = i
        elif textos[0].startswith("Table D") and any(t.startswith("Total") for t in textos[1:]):
            idx["total_d"] = i
    faltan = {"cab_c", "cab_d", "total_c", "total_d"} - idx.keys()
    if faltan:
        raise RecursosNoDisponibles(f"La plantilla no tiene la estructura esperada (faltan {', '.join(sorted(faltan))}).")
    return idx


def _num(v) -> str:
    return f"{v.normalize():f}".replace(".", ",") if v is not None else ""


def rellenar(documento, acta, filas_c, filas_d):
    t, idx = documento.tables[0], localizar_filas(documento)
    for nombre, filas, cab, total in (("C", filas_c, "cab_c", "total_c"), ("D", filas_d, "cab_d", "total_d")):
        libres = idx[total] - idx[cab] - 1
        if len(filas) > libres:
            raise ValueError(f"La plantilla solo tiene {libres} filas libres en la tabla {nombre} y hay {len(filas)}.")
    for i, f in enumerate(filas_c):
        fila = t.rows[idx["cab_c"] + 1 + i]
        _celda(fila, 1, f.codigo, True, WD_ALIGN_PARAGRAPH.LEFT)
        _celda(fila, 2, f.titulo, True, WD_ALIGN_PARAGRAPH.LEFT)
        _celda(fila, 6, "Yes" if f.superada else "No", True)
        _celda(fila, 10, _num(f.ects), True)
        _celda(fila, 12, f.nota_original, True)
    for i, f in enumerate(filas_d):
        fila = t.rows[idx["cab_d"] + 1 + i]
        _celda(fila, 1, f.codigo, True, WD_ALIGN_PARAGRAPH.LEFT)
        _celda(fila, 2, f.titulo, True, WD_ALIGN_PARAGRAPH.LEFT)
        _celda(fila, 8, _num(f.ects), True)
        _celda(fila, 11, f.nota)
    totales = {"total_c": ((10, 11), sum(f.ects for f in filas_c)), "total_d": ((8, 9, 10), sum(f.ects for f in filas_d))}
    for clave, (columnas, total) in totales.items():
        fila = t.rows[idx[clave]]
        for col in columnas:
            if fila.cells[col].text.strip():
                _celda(fila, col, f"Total: {_num(total) or 0}", True)
    m = acta.movilidad
    inicio = f"{m.fecha_llegada:%d/%m/%Y}" if m.fecha_llegada else "(no consta)"
    fin = f"{m.fecha_fin:%d/%m/%Y}" if m.fecha_fin else "(no consta)"
    for fila in t.rows:
        for celda in fila.cells:
            for p in celda.paragraphs:
                completo = "".join(r.text for r in p.runs)
                if "Start and end dates" in completo and p.runs:
                    p.runs[0].text = re.sub(r"from.*?to.*?(?=\n|$)", f"from {inicio} to {fin}", completo, flags=re.DOTALL)
                    for r in p.runs[1:]:
                        r.text = ""
    # quitar las filas vacías sobrantes para que quepa en una página
    borrar = list(range(idx["cab_c"] + 1 + len(filas_c), idx["total_c"]))
    borrar += list(range(idx["cab_d"] + 1 + len(filas_d), idx["total_d"]))
    for tr in [t.rows[i]._tr for i in borrar]:
        tr.getparent().remove(tr)


def construir_docx(acta) -> bytes:
    m = acta.movilidad
    documento = docx.Document(str(recurso("plantilla_simplificado.docx")))
    rellenar(documento, acta, list(acta.filas_c.all()), list(acta.filas_d.all()))
    buffer = io.BytesIO()
    documento.save(buffer)

    entrada = zipfile.ZipFile(io.BytesIO(buffer.getvalue()))
    xml = entrada.read("word/document.xml").decode("utf-8")
    firma = FIRMA_XML.format(firmante=escape(settings.ACTA_FIRMANTE),
                             cargo=escape(settings.ACTA_CARGO_FIRMANTE))
    cola = COLA_PLANTILLA.search(xml)
    if cola:
        xml = xml[:cola.start()] + firma + xml[cola.end():]
    else:
        sect = xml.rfind("<w:sectPr")
        if sect == -1 or sect < xml.rfind("</w:tbl>"):
            raise RecursosNoDisponibles("No encuentro dónde insertar la firma en la plantilla.")
        xml = xml[:sect] + firma + xml[sect:]

    nombre = escape(nombre_para_acta(m.estudiante.nombre_completo))
    cabecera = recurso("header1_template.xml").read_text(encoding="utf-8").replace("__STUDENT_NAME__", nombre)
    cabecera = reemplazar_texto_partido(cabecera, r"\d{4}/\d{4}", curso_largo(m.convocatoria.curso), "Academic Year")

    salida = io.BytesIO()
    with zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as z:
        for item in entrada.infolist():
            if item.filename in ("word/document.xml", "word/header1.xml", "word/_rels/header1.xml.rels",
                                 "word/media/image1.png"):
                continue
            z.writestr(item, entrada.read(item.filename))
        z.writestr("word/document.xml", xml)
        z.writestr("word/header1.xml", cabecera)
        z.writestr("word/media/image1.png", recurso("erasmus_logo.png").read_bytes())
        z.writestr(
            "word/_rels/header1.xml.rels",
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" '
            'Target="media/image1.png"/></Relationships>',
        )
    return salida.getvalue()


def nombre_para_acta(nombre_completo: str) -> str:
    from movilidad.envios import nombre_natural

    return nombre_natural(nombre_completo)


def a_pdf(contenido_docx: bytes) -> bytes | None:
    """PDF con LibreOffice si está instalado; si no, None."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    with tempfile.TemporaryDirectory() as tmp:
        origen = Path(tmp) / "acta.docx"
        origen.write_bytes(contenido_docx)
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp, str(origen)],
                       check=True, capture_output=True, timeout=120)
        return (Path(tmp) / "acta.pdf").read_bytes()


def contar_paginas(pdf: bytes) -> int | None:
    return len(re.findall(rb"/Type\s*/Page[^s]", pdf)) or None


def generar_documentos(acta) -> list:
    """Genera docx (y PDF si se puede) del borrador y devuelve avisos de maquetación."""
    if not acta.editable:
        raise ValueError("El acta ya está firmada: su documento no se regenera.")
    contenido = construir_docx(acta)
    base = f"after_the_mobility{''.join(c for c in nombre_para_acta(acta.movilidad.estudiante.nombre_completo) if c.isalnum())}_DAC_borrador_v{acta.numero}"
    acta.docx.save(f"{base}.docx", ContentFile(contenido), save=False)
    avisos = []
    pdf = a_pdf(contenido)
    if pdf is None:
        avisos.append("No hay LibreOffice: solo se ha generado el docx y no se ha comprobado que quepa en una página.")
    else:
        acta.pdf.save(f"{base}.pdf", ContentFile(pdf), save=False)
        paginas = contar_paginas(pdf)
        if paginas and paginas != 1:
            avisos.append(f"El acta ocupa {paginas} páginas (debería caber en una).")
    acta.avisos = [a for a in acta.avisos if not a.startswith(("No hay LibreOffice", "El acta ocupa"))] + avisos
    acta.save()
    return avisos
