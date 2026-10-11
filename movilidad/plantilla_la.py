"""RF-11: lee la «Plantilla LA.xlsx» del alumno o del coordinador y la compara con el LA vigente.

La plantilla tiene una hoja por fase (Before, During Semestre 1/2, After the mobility) con una tabla
«Asignatura UDC | Créditos | Asignatura Destino | Créditos». Lo que se ha visto en los Excel reales:

- Un «2 por 1» se escribe con la celda UDC combinada (o vacía) en las filas siguientes; un «1 por 2»,
  con la celda de destino vacía. Una fila sin UDC y sin combinar es «sin reconocimiento».
- La tabla «Cambios realizados» (+/−) de las hojas During no se usa: basta con la tabla completa.
- A veces quedan hojas con las filas de ejemplo de la plantilla. Por eso se compara con el LA la
  hoja que mejor casa con él, y se muestra la puntuación de todas.

No se guarda el fichero: solo se lee para la comparación.
"""

import re
import unicodedata
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

import openpyxl

from movilidad.catalogo_udc import parecido
from movilidad.importers.comun import normalizar_nombre
from movilidad.la_vigente import la_vigente
from movilidad.notas import formatear

CABECERA_UDC = {"asignatura udc", "materia udc"}
CABECERA_DESTINO = {"asignatura destino", "materia destino", "asignatura", "component title", "componente"}
CABECERA_CREDITOS = {"creditos", "ects"}
FIN_TABLA = ("total", "cambios realizados")
UMBRAL = 0.7


class PlantillaIlegible(ValueError):
    pass


@dataclass
class GrupoPlantilla:
    destino: list = field(default_factory=list)   # [(nombre, ects | None)]
    udc: list = field(default_factory=list)
    sin_reconocimiento: bool = False


@dataclass
class HojaPlantilla:
    nombre: str
    grupos: list


@dataclass
class Comparacion:
    hoja: HojaPlantilla | None
    puntuaciones: list   # [(nombre de hoja, materias que casan con el LA, materias en la hoja)]
    diferencias: list
    coincidencias: int = 0


def _texto(valor) -> str:
    return "" if valor is None else str(valor).strip()


def _ects(valor):
    try:
        return Decimal(str(valor).replace(",", ".")) if _texto(valor) else None
    except InvalidOperation:
        return None


def _tablas(ws):
    filas = [list(r) for r in ws.iter_rows(values_only=True)]
    combinadas = {(f - 1, c - 1) for rango in ws.merged_cells.ranges
                  for f in range(rango.min_row + 1, rango.max_row + 1)
                  for c in range(rango.min_col, rango.max_col + 1)}
    for i, fila in enumerate(filas):
        norm = [normalizar_nombre(_texto(v)) for v in fila]
        cu = next((k for k, v in enumerate(norm) if v in CABECERA_UDC), None)
        cd = next((k for k, v in enumerate(norm) if v in CABECERA_DESTINO and k != cu), None)
        if cu is None or cd is None or "accion" in norm:
            continue  # la tabla «Cambios realizados» (con columnas «Acción») no es la de correspondencias
        eu = cu + 1 if cu + 1 < len(norm) and norm[cu + 1] in CABECERA_CREDITOS else None
        ed = cd + 1 if cd + 1 < len(norm) and norm[cd + 1] in CABECERA_CREDITOS else None
        entradas, blancos = [], 0
        for j in range(i + 1, len(filas)):
            g = filas[j] + [None] * 30
            u, d = _texto(g[cu]), _texto(g[cd])
            if not u and not d:
                blancos += 1
                if blancos >= 4:
                    break
                continue
            blancos = 0
            if normalizar_nombre(u).startswith(FIN_TABLA) or normalizar_nombre(d).startswith(FIN_TABLA):
                break
            entradas.append({
                "udc": u, "destino": d,
                "ects_udc": _ects(g[eu]) if eu is not None else None,
                "ects_destino": _ects(g[ed]) if ed is not None else None,
                "udc_combinada": (j, cu) in combinadas,
            })
        yield entradas


def _sin_reconocimiento(texto: str) -> bool:
    """«Curso de idiomas (sin reconocimiento)», «sen recoñecemento»… (normalizar_nombre quita los paréntesis)."""
    t = "".join(c for c in unicodedata.normalize("NFKD", texto.lower()) if not unicodedata.combining(c))
    return "reconoc" in t or "reconec" in t


def _agrupar(entradas) -> list:
    grupos = []
    for e in entradas:
        sin_rec = _sin_reconocimiento(e["udc"])
        udc = "" if sin_rec else e["udc"]
        if grupos and e["udc"] and not e["destino"]:            # «1 por 2»: otra UDC del grupo anterior
            grupos[-1].udc.append((udc, e["ects_udc"]))
            continue
        if grupos and not e["udc"] and e["udc_combinada"]:      # «2 por 1»: celda UDC combinada
            grupos[-1].destino.append((e["destino"], e["ects_destino"]))
            continue
        g = GrupoPlantilla(sin_reconocimiento=not udc)
        if e["destino"]:
            g.destino.append((e["destino"], e["ects_destino"]))
        if udc:
            g.udc.append((udc, e["ects_udc"]))
        grupos.append(g)
    # La misma UDC repetida en varias filas (2 por 1 sin combinar) es un único grupo.
    unidos = []
    for g in grupos:
        previo = next((u for u in unidos if g.udc and {normalizar_nombre(n) for n, _ in u.udc}
                       & {normalizar_nombre(n) for n, _ in g.udc}), None)
        if previo:
            previo.destino += [d for d in g.destino if d not in previo.destino]
            previo.udc += [x for x in g.udc if x not in previo.udc]
        else:
            unidos.append(g)
    return unidos


def leer_plantilla(fichero) -> list:
    try:
        wb = openpyxl.load_workbook(fichero, data_only=True)
    except Exception as e:  # noqa: BLE001 (openpyxl lanza de todo con un fichero que no es xlsx)
        raise PlantillaIlegible(f"No se puede leer el Excel: {e}") from e
    hojas = []
    for ws in wb.worksheets:
        for entradas in _tablas(ws):
            if entradas:
                hojas.append(HojaPlantilla(ws.title, _agrupar(entradas)))
    if not hojas:
        raise PlantillaIlegible("No hay ninguna tabla «Asignatura UDC / Asignatura Destino» en el Excel.")
    return hojas


def _casar(nombre, materias):
    """La materia del LA con el nombre más parecido (sin el código que a veces va delante), o None."""
    limpio = re.sub(r"^[A-Z0-9][A-Z0-9\-\.]{3,}\s+", "", nombre.strip())
    puntuadas = sorted(((max(parecido(limpio, a.nombre), parecido(nombre, a.nombre)), a) for a in materias),
                       key=lambda t: -t[0])
    return puntuadas[0][1] if puntuadas and puntuadas[0][0] >= UMBRAL else None


def _relaciones(hoja, destino_la, udc_la, dadas_de_baja=()):
    """Pares (destino, udc) y materias sin reconocimiento de la hoja, con las materias del LA."""
    pares, sin_rec, sueltas, ects = set(), set(), [], []

    def casar(nombre, e, materias, lado):
        a = _casar(nombre, materias)
        if a is None:
            baja = _casar(nombre, dadas_de_baja)
            sueltas.append(f"«{nombre}» se dio de baja en una modificación del LA, pero sigue en el Excel"
                           if baja else f"«{nombre}» ({lado}) no está en el LA vigente")
        elif e is not None and e != a.ects:
            ects.append(f"«{a.nombre}»: {formatear(e)} ECTS en el Excel y {formatear(a.ects)} en el LA")
        return a

    for g in hoja.grupos:
        ds = [a for a in (casar(n, e, destino_la, "destino") for n, e in g.destino) if a]
        us = [a for a in (casar(n, e, udc_la, "UDC") for n, e in g.udc) if a]
        if g.sin_reconocimiento:
            sin_rec.update(ds)
        pares.update((d, u) for d in ds for u in us)
    return pares, sin_rec, sueltas, ects


def comparar(movilidad, hojas) -> Comparacion:
    la = la_vigente(movilidad)
    destino_la = list({a.pk: a for g in la.grupos for a in g.materias_destino}.values())
    udc_la = list({a.pk: a for g in la.grupos for a in g.materias_udc}.values())
    if not la.grupos:
        return Comparacion(None, [], ["No hay LA vigente (ninguna versión firmada) con el que comparar."])

    puntuaciones, mejor = [], None
    for hoja in hojas:
        total = sum(len(g.destino) + len(g.udc) for g in hoja.grupos)
        casan = total - len(_relaciones(hoja, destino_la, udc_la)[2])
        puntuaciones.append((hoja.nombre, casan, total))
        if mejor is None or casan >= mejor[1]:  # a igualdad, la hoja posterior (refleja los cambios)
            mejor = (hoja, casan)
    hoja = mejor[0]
    bajas = list({(type(b.asignatura), b.asignatura.pk): b.asignatura for b in la.bajas}.values())
    pares, sin_rec, sueltas, ects = _relaciones(hoja, destino_la, udc_la, bajas)

    pares_la = {(d, u) for g in la.grupos for d in g.materias_destino for u in g.materias_udc}
    sin_rec_la = {d for g in la.grupos if g.sin_reconocimiento for d in g.materias_destino}
    en_excel = {d for d, _ in pares} | sin_rec | {u for _, u in pares}

    diferencias = [f"{s}." for s in sueltas]
    diferencias += [f"«{a.nombre}» está en el LA vigente pero no en el Excel."
                    for a in destino_la + udc_la if a not in en_excel]
    for d, u in sorted(pares - pares_la, key=lambda p: p[0].nombre):
        diferencias.append(f"En el Excel «{d.nombre}» reconoce «{u.nombre}»; en el LA vigente, no.")
    for d, u in sorted(pares_la - pares, key=lambda p: p[0].nombre):
        if d in en_excel and u in en_excel:
            diferencias.append(f"En el LA vigente «{d.nombre}» reconoce «{u.nombre}»; en el Excel, no.")
    diferencias += [f"«{d.nombre}» va sin reconocimiento en el Excel pero no en el LA." for d in sin_rec - sin_rec_la]
    diferencias += [f"«{d.nombre}» va sin reconocimiento en el LA pero no en el Excel."
                    for d in sin_rec_la - sin_rec if d in en_excel]
    diferencias += sorted(set(ects))
    return Comparacion(hoja, puntuaciones, diferencias, coincidencias=len(pares & pares_la))
