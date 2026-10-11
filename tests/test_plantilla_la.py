"""RF-11: comparar la «Plantilla LA.xlsx» con el LA vigente (Excel sintéticos generados aquí)."""

import io
from decimal import Decimal

import openpyxl
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile

from movilidad.models import AsignaturaUDC, TipoVersionLA
from movilidad.plantilla_la import PlantillaIlegible, comparar, leer_plantilla
from tests.test_acuerdo_academico import A, B, dest, gei, mov, version  # noqa: F401 (fixtures)

pytestmark = pytest.mark.django_db
CABECERA = ["Asignatura UDC", "Créditos", "Asignatura Destino", "Créditos"]


def excel(hojas):
    """hojas: {nombre: [filas]}; una fila «*» combina la celda UDC con la de encima."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nombre, filas in hojas.items():
        ws = wb.create_sheet(nombre)
        ws.append([f"Asignaturas ({nombre})"])
        ws.append(CABECERA)
        for fila in filas:
            if fila[0] == "*":
                ws.append([None, None, *fila[1:]])
                ws.merge_cells(start_row=ws.max_row - 1, start_column=1, end_row=ws.max_row, end_column=1)
            else:
                ws.append(list(fila))
        ws.append(["Total", 30, "Total", 30])
        ws.append([])
        ws.append(["Cambios realizados (+ es Añadir, - es Eliminar)"])
        ws.append(["Asignatura UDC", "Acción", "Asignatura destino", "Acción"])
        ws.append(["Xestión de Proxectos", "-", "Project Management", "-"])
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


@pytest.fixture
def la(mov, gei):
    def u(codigo, nombre):
        return AsignaturaUDC.objects.create(curso_academico="2026/27", titulacion=gei, codigo=codigo, nombre=nombre,
                                            ects=Decimal(6))
    xi, xp, ap = u("200190", "Xestión de Infraestruturas"), u("202323", "Xestión de Proxectos"), u("202376", "Algoritmos")
    d = {n: dest(mov, n, e) for n, e in (("Cloud Computing", 6), ("Project Management", 6), ("Algorithms I", 3),
                                          ("Algorithms II", 3), ("Lithuanian A1", 3), ("Software Quality", 6))}
    version(mov, filas=[(A, 1, [d["Cloud Computing"]], [xi]), (A, 1, [d["Project Management"]], [xp]),
                        (A, 1, [d["Algorithms I"], d["Algorithms II"]], [ap]), (A, 1, [d["Lithuanian A1"]], [], True)])
    # During: cambia Project Management por Software Quality
    version(mov, tipo=TipoVersionLA.DURING_S1, filas=[(B, 1, [d["Project Management"]], [xp]),
                                                      (A, 1, [d["Software Quality"]], [xp])])
    return d


EJEMPLO_PLANTILLA = [("Bases de datos", 6, "Databases 1", 5), ("Sistemas operativos", 6, "Operating systems", 7)]
DURING_BIEN = [
    ("Gestión de Infraestructuras", 6, "Cloud Computing", 6),       # en castellano: casa con el gallego
    ("Gestión de Proyectos", 6, "Software Quality", 6),
    ("Algoritmos", 6, "Algorithms I", 3),
    ("*", "Algorithms II", 3),                                      # 2 por 1 con la celda UDC combinada
    ("Curso de idiomas (sin reconocimiento)", 3, "Lithuanian A1", 3),
]


def test_elige_la_hoja_que_casa_y_no_ve_diferencias(mov, la):
    hojas = leer_plantilla(excel({"Before": DURING_BIEN[:1] + [("Gestión de Proyectos", 6, "Project Management", 6)]
                                  + DURING_BIEN[2:], "During Semestre 1": DURING_BIEN,
                                  "After the mobility": EJEMPLO_PLANTILLA}))
    assert [h.nombre for h in hojas] == ["Before", "During Semestre 1", "After the mobility"]  # sin «Cambios»
    c = comparar(mov, hojas)
    assert c.hoja.nombre == "During Semestre 1"
    assert c.diferencias == [] and c.coincidencias == 4
    assert ("After the mobility", 0, 4) in c.puntuaciones  # filas de ejemplo de la plantilla


def test_senala_las_diferencias(mov, la):
    filas = [
        ("Gestión de Infraestructuras", 6, "Cloud Computing", 5),     # ECTS distintos
        ("Gestión de Proyectos", 6, "Project Management", 6),         # el Excel no recoge el During
        ("Algoritmos", 6, "Algorithms I", 3),
        ("Algoritmos", 6, "Algorithms II", 3),                        # 2 por 1 repitiendo la UDC: es lo mismo
        ("", None, "Lithuanian A1", 3),                               # sin UDC y sin combinar: sin reconocimiento
        ("Robótica", 6, "Underwater Robotics", 6),
    ]
    c = comparar(mov, leer_plantilla(excel({"During Semestre 1": filas})))
    d = " | ".join(c.diferencias)
    assert "«Robótica» (UDC) no está en el LA vigente" in d and "«Underwater Robotics» (destino)" in d
    assert "«Software Quality» está en el LA vigente pero no en el Excel" in d
    assert "«Project Management» se dio de baja en una modificación del LA, pero sigue en el Excel" in d
    assert "«Cloud Computing»: 5 ECTS en el Excel y 6 en el LA" in d
    assert "Algorithms" not in d and "Lithuanian" not in d
    assert c.coincidencias == 3


def test_sin_la_vigente_o_excel_raro(mov):
    c = comparar(mov, leer_plantilla(excel({"Before": EJEMPLO_PLANTILLA})))
    assert c.hoja is None and "No hay LA vigente" in c.diferencias[0]
    with pytest.raises(PlantillaIlegible, match="No se puede leer"):
        leer_plantilla(io.BytesIO(b"esto no es un xlsx"))
    wb = openpyxl.Workbook()
    wb.active.append(["Nada", "que", "ver"])
    vacio = io.BytesIO()
    wb.save(vacio)
    vacio.seek(0)
    with pytest.raises(PlantillaIlegible, match="No hay ninguna tabla"):
        leer_plantilla(vacio)


def test_vista_del_admin(admin_client, mov, la):
    url = f"/admin/movilidad/movilidad/{mov.pk}/comparar-plantilla/"
    assert "Comparar con la plantilla" in admin_client.get(f"/admin/movilidad/movilidad/{mov.pk}/change/").content.decode()
    assert admin_client.get(url).status_code == 200
    fichero = SimpleUploadedFile("Plantilla LA.xlsx", excel({"During Semestre 1": DURING_BIEN}).read())
    r = admin_client.post(url, {"excel": fichero})
    assert "El Excel coincide con el LA vigente" in r.content.decode()
    r = admin_client.post(url, {"excel": SimpleUploadedFile("x.xlsx", b"nada")})
    assert "No se puede leer el Excel" in r.content.decode()
