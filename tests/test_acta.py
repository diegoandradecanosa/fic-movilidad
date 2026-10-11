"""Tests del bloque 4: escalas y conversión de notas, ToR, acta (tablas C/D, versiones), docx,
envío del «ATM de …» y extracción con IA. Todo sintético: escalas, plantilla y PDF inventados."""

import io
import json
import zipfile
from datetime import date, timedelta
from decimal import Decimal

import docx
import pytest
from django.core import mail
from django.core.files.base import ContentFile
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.db import transaction

from movilidad import extraccion_ia
from movilidad.acta import acta_vigente, crear_correccion, generar_borrador, proponer_emparejamiento
from movilidad.acta_docx import construir_docx, curso_largo, generar_documentos, reemplazar_texto_partido
from movilidad.envios import EnvioImposible, enviar_atm
from movilidad.importers.equivalencias import importar_equivalencias
from movilidad.models import (
    Acta,
    EscalaNotas,
    EstadoActa,
    EstadoMovilidad,
    EstadoPlazo,
    EstadoToR,
    EventoMovilidad,
    LineaToR,
    ModificacionNoPermitida,
    SemestresToR,
    Superada,
    TipoPlazo,
    ToR,
    VersionLA,
)
from movilidad.notas import NO_PRESENTADO, SUSPENSO, combinar, convertir, elegir_escala, formatear, resultado_linea
from tests.test_acuerdo_academico import A, dest, gei, mov, udc, version  # noqa: F401 (fixtures)

pytestmark = pytest.mark.django_db

# Escala inventada de 18 a 30 (con «30L»), solo notas de aprobado, como las tablas oficiales.
EQ_18_30 = {**{str(n): v for n, v in zip(range(18, 31), (5, 5.4, 5.8, 6.2, 6.6, 7, 7.4, 7.8, 8.2, 8.6, 9, 9.4, 9.8))},
            "30L": 10}


def eq(tabla):
    return {k: {"equivalente_5_10": v, "equivalente_1_4": 1, "equivalente_literal": "X"} for k, v in tabla.items()}


def escala(pais_iso, id_escala, tabla, pais="FICTICIA", minima="", maxima=""):
    return EscalaNotas.objects.create(pais=pais, pais_iso=pais_iso, id_escala=id_escala, calificaciones=eq(tabla),
                                      num_calificaciones=len(tabla), nota_minima_aprobatoria=minima, nota_maxima=maxima)


@pytest.fixture
def escala_pt():
    return escala("PT", "18(1ª)-30L(14ª)", EQ_18_30, minima="18", maxima="30")


# --- conversión de notas ---------------------------------------------------------------------------

def test_convertir_y_formatear(escala_pt):
    assert convertir(escala_pt, "30L")["equivalente_5_10"] == 10
    assert convertir(escala_pt, "30")["equivalente_5_10"] == 9.8
    assert convertir(escala_pt, "27,0")["equivalente_5_10"] == 8.6
    assert convertir(escala_pt, "15") is None
    assert (formatear(8.0), formatear(8.5), formatear(6.234)) == ("8", "8,5", "6,23")


def test_elegir_escala_y_ambiguedad():
    fina = escala("SE", "VG(2)", {"G": 6.5, "VG": 8.5}, pais="SUECIA")
    escala("SE", "E-A", {"E": 5, "D": 6, "C": 7, "B": 8, "A": 9.5, "VG": 8.0, "G": 6.0}, pais="SUECIA")
    eleccion = elegir_escala(EscalaNotas.objects.filter(pais_iso="SE"), ["VG", "G"])
    assert eleccion.escala == fina  # la más sencilla que cubre las notas
    assert eleccion.ambigua and "E-A" in eleccion.motivo  # la otra daría 8 en vez de 8,5
    assert not elegir_escala(EscalaNotas.objects.filter(pais_iso="SE"), ["A", "B"]).ambigua


def test_resultado_de_una_materia(escala_pt):
    assert resultado_linea(escala_pt, "27").nota == "8,6"
    # suspenso numérico: equivalente proporcional al mínimo de aprobado (5 × 15 / 18 = 4,17)
    assert resultado_linea(escala_pt, "15").nota == "Suspenso (4,17)"
    assert resultado_linea(escala_pt, "0").nota == "Suspenso (0)"
    assert resultado_linea(escala_pt, "F").nota == SUSPENSO  # letra: sin equivalente, se completa a mano
    assert resultado_linea(escala_pt, "ABS").nota == NO_PRESENTADO
    assert resultado_linea(escala_pt, "").nota == NO_PRESENTADO
    # el ToR la da por no superada aunque la escala la aprueba: sin número, para revisar
    assert resultado_linea(escala_pt, "27", superada_declarada=False).nota == SUSPENSO
    # escala en la que menos es mejor (1 = sobresaliente): no se inventa el equivalente
    invertida = escala("DE", "1,0-4,0", {"1,0": 10, "2,0": 8, "3,0": 6.5, "4,0": 5})
    assert resultado_linea(invertida, "5,0").nota == SUSPENSO
    revisar = resultado_linea(escala_pt, "P", superada_declarada=True)
    assert revisar.nota == "" and revisar.superada and "revisar" in revisar.detalle


def test_dos_por_uno_media_ponderada(escala_pt):
    # misma escala: se pondera la nota original (27·3 + 30·6)/9 = 29 y se convierte una vez
    r = combinar([(escala_pt, "27", 3, None), (escala_pt, "30", 6, None)])
    assert (r.nota, r.superada) == ("9,4", True)
    # letras (no promediables): se convierte cada una y se pondera (6,5·2 + 8,5·4)/6 = 7,83
    letras = escala("SE", "VG(2)", {"G": 6.5, "VG": 8.5})
    assert combinar([(letras, "G", 2, None), (letras, "VG", 4, None)]).nota == "7,83"
    # si una no está superada, la UDC no se reconoce
    assert combinar([(escala_pt, "27", 3, None), (escala_pt, "ABS", 3, None)]).nota == NO_PRESENTADO
    assert combinar([(escala_pt, "27", 3, None), (escala_pt, "12", 3, None)]).nota == SUSPENSO


def test_importar_equivalencias(tmp_path):
    datos = {"metadata": {}, "sistema_espanol": {}, "paises": {
        "ITALIA": {"escalas": [{"id_escala": "18-30L", "num_calificaciones": 2, "nota_minima_aprobatoria": "18",
                                "nota_maxima": "30L", "calificaciones": eq({"18": 5, "30L": 10})}]},
        "REPUBLICA DE COREA": {"escalas": [{"id_escala": "F-A+", "calificaciones": eq({"A+": 10})}]},
        "ATLANTIDA": {"escalas": [{"id_escala": "x", "calificaciones": {}}]},
    }}
    ruta = tmp_path / "equivalencias.json"
    ruta.write_text(json.dumps(datos), encoding="utf-8")
    r = importar_equivalencias(ruta)
    assert (r["paises"], r["escalas"], r["sin_iso"]) == (3, 3, ["ATLANTIDA"])
    assert EscalaNotas.objects.get(pais="ITALIA").pais_iso == "IT"
    assert EscalaNotas.objects.get(pais="REPUBLICA DE COREA").pais_iso == "KR"
    call_command("importar_equivalencias", str(ruta))
    assert EscalaNotas.objects.count() == 3


# --- ToR ------------------------------------------------------------------------------------------

def pdf(nombre="doc.pdf", contenido=b"%PDF-1.4 sintetico"):
    return SimpleUploadedFile(nombre, contenido, content_type="application/pdf")


@pytest.fixture
def media(settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path / "media"
    return settings


def nuevo_tor(mov, lineas, semestres=SemestresToR.S1, aceptar=True, **kw):
    """lineas: (titulo, ects, nota, asignatura_destino[, superada])"""
    tor = ToR.objects.create(movilidad=mov, fecha_recepcion=kw.pop("fecha", date(2027, 2, 1)), semestres=semestres, **kw)
    for i, (titulo, ects, nota, asig, *sup) in enumerate(lineas):
        LineaToR.objects.create(tor=tor, orden=i, titulo=titulo, ects=Decimal(ects), nota=nota, asignatura_destino=asig,
                                superada=sup[0] if sup else Superada.DESCONOCIDO)
    if aceptar:
        tor.aceptar()
    return tor


def test_registrar_tor_avanza_y_crea_el_plazo_de_cinco_semanas(mov):
    mov.estado = EstadoMovilidad.FINALIZADA
    mov.save()
    tor = nuevo_tor(mov, [], fecha=date(2027, 2, 1))
    mov.refresh_from_db()
    assert mov.estado == EstadoMovilidad.TOR_RECIBIDO
    plazo = mov.plazos.get(tipo=TipoPlazo.ENVIAR_ACTA)
    assert (plazo.tor, plazo.fecha_limite, plazo.regla) == (tor, date(2027, 3, 8), "RN-15")


def test_nombre_exacto_gana_aunque_haya_otro_parecido(mov, gei):
    a1, a2 = dest(mov, "Algorithms I"), dest(mov, "Algorithms II")
    version(mov, filas=[(A, 1, [a1, a2], [udc(gei, "200950")])])
    tor = nuevo_tor(mov, [("Algorithms I", 5, "1,7", None), ("ALGORITHMS II", 5, "2,3", None)], aceptar=False)
    assert proponer_emparejamiento(tor) == []
    assert [l.asignatura_destino for l in tor.lineas.all()] == [a1, a2]


def test_los_suspensos_no_vuelven_ambigua_la_escala(mov, gei, escala_pt):
    d1, d2 = dest(mov, "X"), dest(mov, "Y")
    version(mov, filas=[(A, 1, [d1], [udc(gei, "200960")]), (A, 1, [d2], [udc(gei, "200961")])])
    nuevo_tor(mov, [("X", 6, "27", d1), ("Y", 6, "12", d2)])  # 12: suspenso, no está en ninguna escala
    acta = generar_borrador(mov)
    assert not any("ambigua" in a for a in acta.avisos)
    assert [f.nota for f in acta.filas_d.all()] == ["8,6", "Suspenso (3,33)"]


def test_lineas_de_un_tor_aceptado_no_se_tocan(mov):
    tor = nuevo_tor(mov, [("Redes", 6, "27", None)])
    linea = tor.lineas.get()
    linea.nota = "30"
    with pytest.raises(ModificacionNoPermitida):
        linea.save()


def test_proponer_emparejamiento(mov, gei):
    redes, ia, ia2 = dest(mov, "Computer Networks"), dest(mov, "Artificial Intelligence I"), dest(mov, "Artificial Intelligence II")
    redes.codigo = "CN101"
    redes.save()
    version(mov, filas=[(A, 1, [redes], [udc(gei, "200901")]), (A, 1, [ia], [udc(gei, "200902")]),
                        (A, 1, [ia2], [udc(gei, "200903")])])
    tor = nuevo_tor(mov, [("Networks (lab)", 6, "27", None), ("Artificial Intelligence 1", 6, "28", None),
                          ("Underwater Basket Weaving", 3, "30", None)], aceptar=False)
    tor.lineas.filter(titulo__startswith="Networks").update(codigo="cn-101")
    sin_pareja = proponer_emparejamiento(tor)
    asignadas = {l.titulo: l.asignatura_destino for l in tor.lineas.all()}
    assert asignadas["Networks (lab)"] == redes  # por código
    assert asignadas["Artificial Intelligence 1"] is None  # I y II se parecen demasiado: no adivina
    assert {l.titulo for l in sin_pareja} == {"Artificial Intelligence 1", "Underwater Basket Weaving"}


# --- borrador del acta ----------------------------------------------------------------------------

@pytest.fixture
def caso(mov, gei, escala_pt):
    """LA con 1:1, «2 por 1», 1 → 2 UDC, curso de idioma sin reconocimiento y una materia suspensa."""
    d = {n: dest(mov, n, e) for n, e in (("Databases", 6), ("Algorithms I", 3), ("Algorithms II", 6),
                                          ("Machine Learning", 6), ("Portuguese A2", 3), ("Compilers", 6),
                                          ("Robotics", 6))}
    u = {c: udc(gei, c, e) for c, e in (("200001", 6), ("200002", 9), ("200003", 6), ("200004", 6),
                                         ("200005", 6), ("200006", 6))}
    u["200001"].codigo_oficial = "614G01001"
    u["200001"].save()
    version(mov, filas=[
        (A, 1, [d["Databases"]], [u["200001"]]),
        (A, 1, [d["Algorithms I"], d["Algorithms II"]], [u["200002"]]),
        (A, 1, [d["Machine Learning"]], [u["200003"], u["200004"]]),
        (A, 1, [d["Portuguese A2"]], [], True),
        (A, 1, [d["Compilers"]], [u["200005"]]),
        (A, 2, [d["Robotics"]], [u["200006"]]),
    ])
    mov.estado = EstadoMovilidad.TOR_RECIBIDO
    mov.save()
    tor = nuevo_tor(mov, [
        ("Databases", 6, "24", d["Databases"]),
        ("Algorithms I", 3, "27", d["Algorithms I"]),
        ("Algorithms II", 6, "30", d["Algorithms II"]),
        ("Machine Learning", 6, "28", d["Machine Learning"]),
        ("Portuguese A2", 3, "30L", d["Portuguese A2"]),
        ("Compilers", 6, "ABS", d["Compilers"]),
        ("Surf", 2, "30", None),
    ], nombre_en_documento="EXEMPLO, UXIA")
    return {"d": d, "u": u, "tor": tor, "mov": mov}


def test_generar_borrador_tablas_c_y_d(caso):
    acta = generar_borrador(caso["mov"])
    c = [(f.titulo, f.superada, f.ects, f.nota_original) for f in acta.filas_c.all()]
    assert len(c) == 7
    assert ("Compilers", False, Decimal(6), "ABS") in c  # no superada: ECTS nominales y nota literal
    assert ("Surf", True, Decimal(2), "30") in c  # fuera del LA: solo tabla C
    d = {f.titulo: (f.codigo, f.ects, f.nota) for f in acta.filas_d.all()}
    assert d["UDC 200001"] == ("614G01001", Decimal(6), "7,4")
    assert d["UDC 200002"] == ("200002", Decimal(9), "9,4")  # (27·3 + 30·6)/9 = 29 → 9,4
    assert d["UDC 200003"][2] == d["UDC 200004"][2] == "9"  # una de destino cubre dos UDC
    assert d["UDC 200005"] == ("200005", Decimal(0), NO_PRESENTADO)
    assert "UDC 200006" not in d  # 2.º semestre: el ToR es solo del 1.º
    assert len(acta.filas_d.all()) == 5
    avisos = " | ".join(acta.avisos)
    assert "«Surf» no está en el LA vigente" in avisos
    assert "Pendiente del ToR del 2.º semestre" in avisos
    assert "no coincide" not in avisos  # «EXEMPLO» coincide con el apellido
    caso["mov"].refresh_from_db()
    assert caso["mov"].estado == EstadoMovilidad.ACTA_BORRADOR


def test_materia_del_la_que_no_esta_en_el_tor_es_no_presentado(caso):
    linea = caso["tor"].lineas.get(titulo="Databases")
    caso["tor"].estado = EstadoToR.BORRADOR
    caso["tor"].save()
    linea.delete()
    caso["tor"].aceptar()
    acta = generar_borrador(caso["mov"])
    d = {f.titulo: (f.ects, f.nota) for f in acta.filas_d.all()}
    assert d["UDC 200001"] == (Decimal(0), NO_PRESENTADO)
    assert any("va como «No presentado»" in a for a in acta.avisos)
    assert "UDC 200006" not in d  # el 2.º semestre sigue pendiente, no es «No presentado»


def test_solo_usa_tor_aceptados_y_regenera_el_mismo_borrador(caso):
    caso["tor"].estado = EstadoToR.BORRADOR
    caso["tor"].save()
    with pytest.raises(ValueError, match="aceptado"):
        generar_borrador(caso["mov"])
    caso["tor"].aceptar()
    primera = generar_borrador(caso["mov"])
    segunda = generar_borrador(caso["mov"])
    assert primera.pk == segunda.pk and Acta.objects.count() == 1 and segunda.filas_c.count() == 7


def test_avisos_de_nombre_y_de_escala(mov, gei):
    d = dest(mov, "X")
    version(mov, filas=[(A, 1, [d], [udc(gei, "200100")])])
    nuevo_tor(mov, [("X", 6, "8", d)], nombre_en_documento="Outra Persoa, Ana")
    avisos = " | ".join(generar_borrador(mov).avisos)
    assert "no coincide" in avisos and "no hay escalas de notas para el país PT" in avisos


def test_firmar_corregir_y_acta_vigente(caso, media):
    acta = generar_borrador(caso["mov"])
    acta.firmar(pdf("acta_firmada.pdf"), date(2027, 2, 20))
    caso["mov"].refresh_from_db()
    assert caso["mov"].estado == EstadoMovilidad.ACTA_FIRMADA
    fila = acta.filas_d.first()
    fila.nota = "10"
    with pytest.raises(ModificacionNoPermitida), transaction.atomic():
        fila.save()
    acta.motivo_correccion = "otra cosa"
    with pytest.raises(ModificacionNoPermitida):
        acta.save()
    acta = Acta.objects.get(pk=acta.pk)
    v2 = crear_correccion(acta, "Faltaba la nota de Robotics")
    assert (v2.numero, v2.corrige, v2.estado) == (2, acta, EstadoActa.BORRADOR)
    assert acta_vigente(caso["mov"]) == acta  # hasta que se firme la v2
    v2.firmar(pdf("acta_v2.pdf"), date(2027, 3, 1))
    assert acta_vigente(caso["mov"]) == v2
    assert Acta.objects.get(pk=acta.pk).estado == EstadoActa.FIRMADA  # la v1 no se toca


def test_no_se_corrige_un_borrador(caso):
    with pytest.raises(ValueError):
        crear_correccion(generar_borrador(caso["mov"]), "x")


# --- envío «ATM de …» -----------------------------------------------------------------------------

@pytest.fixture
def correo(media):
    media.CORREO_ADMINISTRACION = "administracion@example.org"
    media.CORREO_COORDINADOR = media.DEFAULT_FROM_EMAIL = "coordinacion@example.org"
    media.CORREO_ORI = "ori@example.org"
    return media


def test_enviar_atm(caso, correo):
    caso["tor"].pdf.save("tor.pdf", ContentFile(b"%PDF tor"))
    VersionLA.objects.filter(movilidad=caso["mov"]).update(pdf="la/la_before.pdf")  # PDF ya en el almacenamiento
    (correo.MEDIA_ROOT / "la").mkdir(parents=True)
    (correo.MEDIA_ROOT / "la" / "la_before.pdf").write_bytes(b"%PDF la")
    acta = generar_borrador(caso["mov"])
    with pytest.raises(EnvioImposible):
        enviar_atm(acta)  # sin firmar
    acta.firmar(pdf("acta_firmada.pdf"), date(2027, 2, 20))
    c = enviar_atm(acta)
    m = mail.outbox[0]
    assert m.subject == "ATM de Uxía Proba Exemplo"
    assert m.to == ["administracion@example.org"] and m.cc == ["ori@example.org", "coordinacion@example.org"]
    assert [a[0][:5] for a in m.attachments] == ["acta_", "tor.p", "la_be"]
    assert c.enviada and c.acta == acta
    acta.refresh_from_db()
    caso["mov"].refresh_from_db()
    assert acta.estado == EstadoActa.ENVIADA and acta.fecha_envio
    assert caso["mov"].estado == EstadoMovilidad.ENVIADA_ADMINISTRACION
    assert caso["mov"].plazos.get(tipo=TipoPlazo.ENVIAR_ACTA).estado == EstadoPlazo.CUMPLIDO
    with pytest.raises(EnvioImposible):
        enviar_atm(acta)  # ya enviada
    # corrección: «RE: ATM de …» y la movilidad vuelve a borrador
    v2 = crear_correccion(acta, "Escala equivocada")
    caso["mov"].refresh_from_db()
    assert caso["mov"].estado == EstadoMovilidad.ACTA_BORRADOR
    v2.firmar(pdf("acta_v2.pdf"), date(2027, 3, 1))
    enviar_atm(v2)
    assert mail.outbox[1].subject == "RE: ATM de Uxía Proba Exemplo"
    assert "Escala equivocada" in mail.outbox[1].body
    estados = list(EventoMovilidad.objects.filter(movilidad=caso["mov"]).order_by("id")
                   .values_list("estado_nuevo", flat=True))
    assert estados[-6:] == ["ACTA_BORRADOR", "ACTA_FIRMADA", "ENVIADA_ADMINISTRACION", "ACTA_BORRADOR",
                            "ACTA_FIRMADA", "ENVIADA_ADMINISTRACION"]


def test_atm_sin_configuracion(caso, media):
    media.CORREO_ADMINISTRACION = ""
    acta = generar_borrador(caso["mov"])
    acta.firmar(pdf("acta.pdf"), date(2027, 2, 20))
    with pytest.raises(EnvioImposible, match="CORREO_ADMINISTRACION"):
        enviar_atm(acta)


# --- docx -----------------------------------------------------------------------------------------

def plantilla_sintetica(ruta):
    """Plantilla con la misma estructura que la «Simplificado»: tabla de 15 columnas con C y D."""
    d = docx.Document()
    t = d.add_table(rows=0, cols=15)

    def fila(textos):
        celdas = t.add_row().cells
        for col, texto in textos.items():
            celdas[col].text = texto

    fila({0: "Start and end dates of the study period: from 01/01/2000 to 02/02/2000"})
    fila({1: "Component code", 2: "Component title at the Receiving Institution", 12: "Grades received"})
    for _ in range(10):
        fila({})
    fila({0: "Table C", 2: "x", 10: "Total: 99", 11: "Total: 99"})
    fila({1: "Component code", 2: "Component title recognised", 11: "Grades registered"})
    for _ in range(10):
        fila({})
    fila({0: "Table D", 2: "x", 8: "Total: 99", 9: "Total: 99", 10: "Total: 99"})
    d.save(ruta / "plantilla_simplificado.docx")
    (ruta / "header1_template.xml").write_text(
        '<w:hdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:p>'
        '<w:r><w:t>__STUDENT_NAME__</w:t></w:r><w:r><w:t>Academic Year 202</w:t></w:r>'
        '<w:r><w:t>5/202</w:t></w:r><w:r><w:t>6</w:t></w:r></w:p></w:hdr>', encoding="utf-8")
    (ruta / "erasmus_logo.png").write_bytes(b"\x89PNG sintetico")


@pytest.fixture
def recursos(settings, tmp_path):
    ruta = tmp_path / "recursos_acta"
    ruta.mkdir()
    plantilla_sintetica(ruta)
    settings.RECURSOS_ACTA_DIR = str(ruta)
    settings.ACTA_FIRMANTE = "Persoa Coordinadora"
    return ruta


def test_reemplazar_texto_partido_y_curso():
    xml = "<a><t>Academic Year 202</t><t>5/202</t><t>6</t></a>"
    assert reemplazar_texto_partido(xml, r"\d{4}/\d{4}", "2026/2027", "Academic Year") == \
        "<a><t>Academic Year 202</t><t>6/202</t><t>7</t></a>"
    doble = xml + xml  # la plantilla repite la cabecera (versión alternativa)
    assert reemplazar_texto_partido(doble, r"\d{4}/\d{4}", "2026/2027", "Academic Year").count("<t>7</t>") == 2
    assert curso_largo("2026/27") == "2026/2027"


def test_construir_docx(caso, recursos):
    acta = generar_borrador(caso["mov"])
    contenido = construir_docx(acta)
    z = zipfile.ZipFile(io.BytesIO(contenido))
    cabecera = z.read("word/header1.xml").decode()
    assert "Uxía Proba Exemplo" in cabecera and "2026/2027" in "".join(cabecera.split("</w:t></w:r><w:r><w:t>"))
    cuerpo = z.read("word/document.xml").decode()
    assert "A Coruña, en la fecha que figura en la firma digital" in cuerpo and "Asdo: Persoa Coordinadora" in cuerpo
    tabla = docx.Document(io.BytesIO(contenido)).tables[0]
    textos = [[c.text for c in f.cells] for f in tabla.rows]
    assert len(textos) == 2 + 7 + 1 + 1 + 5 + 1  # filas vacías sobrantes eliminadas
    assert textos[2][2] == "Databases" and textos[2][12] == "24"
    assert any(f[2] == "UDC 200002" and f[11] == "9,4" for f in textos)
    assert any("Total: 32" in f[10] for f in textos)  # tabla C: 6+3+6+6+3+6+2
    assert "from (no consta) to (no consta)" in textos[0][0]


def test_generar_documentos_sin_libreoffice(caso, recursos, media, monkeypatch):
    monkeypatch.setattr("shutil.which", lambda _: None)
    acta = generar_borrador(caso["mov"])
    avisos = generar_documentos(acta)
    assert acta.docx.name.endswith(".docx") and not acta.pdf
    assert "No hay LibreOffice" in avisos[0]


def test_sin_recursos_avisa(caso, settings, tmp_path):
    from movilidad.acta_docx import RecursosNoDisponibles

    settings.RECURSOS_ACTA_DIR = str(tmp_path / "no_existe")
    with pytest.raises(RecursosNoDisponibles, match="RECURSOS_ACTA_DIR"):
        construir_docx(generar_borrador(caso["mov"]))


# --- extracción con IA ----------------------------------------------------------------------------

RESPUESTA_IA = {
    "nombre_alumno": "Uxía Proba", "descripcion_escala": "18 to 30, 30L with honours", "nota_minima": "18",
    "nota_maxima": "30", "observaciones": "",
    "materias": [{"codigo": "DB1", "titulo": "Databases", "ects": 6, "nota": "24", "superada": "SI"},
                 {"codigo": "", "titulo": "Compilers", "ects": 6, "nota": "ABS", "superada": "NO"}],
}


def test_extraer_tor_con_ia(caso, media):
    tor = ToR.objects.create(movilidad=caso["mov"], fecha_recepcion=date(2027, 2, 1), pdf=pdf("tor.pdf"))
    assert extraccion_ia.extraer_tor(tor, llamar=lambda contenido: RESPUESTA_IA) == 2
    tor.refresh_from_db()
    assert tor.estado == EstadoToR.BORRADOR and tor.extraido_con_ia  # borrador: lo acepta el coordinador
    assert [(l.titulo, l.nota, l.superada) for l in tor.lineas.all()] == [
        ("Databases", "24", Superada.SI), ("Compilers", "ABS", Superada.NO)]
    assert tor.nota_minima == "18" and tor.nombre_en_documento == "Uxía Proba"


def test_ia_desactivada_sin_clave(caso, media, settings):
    settings.ANTHROPIC_API_KEY = ""
    tor = ToR.objects.create(movilidad=caso["mov"], fecha_recepcion=date(2027, 2, 1), pdf=pdf("tor.pdf"))
    assert not extraccion_ia.disponible()
    with pytest.raises(extraccion_ia.IANoDisponible, match="ANTHROPIC_API_KEY"):
        extraccion_ia.extraer_tor(tor)


def test_llamada_a_la_api(settings, monkeypatch):
    """La petición usa el modelo configurado, structured outputs y el PDF como documento."""
    import anthropic

    settings.ANTHROPIC_API_KEY, settings.ANTHROPIC_MODEL = "clave-de-prueba", "claude-opus-5-5"
    llamadas = []

    class Bloque:
        type, text = "text", json.dumps(RESPUESTA_IA)

    class Respuesta:
        stop_reason, content = "end_turn", [Bloque()]

    class Mensajes:
        def create(self, **kw):
            llamadas.append(kw)
            return Respuesta()

    class Cliente:
        def __init__(self, **kw):
            self.beta = type("Beta", (), {"messages": Mensajes()})()

    monkeypatch.setattr(anthropic, "Anthropic", Cliente)
    assert extraccion_ia.llamar_modelo(b"%PDF")["materias"][0]["titulo"] == "Databases"
    kw = llamadas[0]
    assert kw["model"] == "claude-opus-5-5"
    assert kw["output_config"]["format"]["type"] == "json_schema"
    assert kw["messages"][0]["content"][0]["type"] == "document"
    Respuesta.stop_reason = "refusal"
    with pytest.raises(extraccion_ia.IANoDisponible, match="declinó"):
        extraccion_ia.llamar_modelo(b"%PDF")


# --- admin ----------------------------------------------------------------------------------------

def test_admin_tor_y_acta(admin_client, caso, recursos, media, monkeypatch):
    monkeypatch.setattr("shutil.which", lambda _: None)
    tor = caso["tor"]
    assert admin_client.get(f"/admin/movilidad/tor/{tor.pk}/change/").status_code == 200
    r = admin_client.post("/admin/movilidad/tor/", {"action": "generar_acta", "_selected_action": [tor.pk]})
    assert r.status_code == 302
    acta = Acta.objects.get()
    r = admin_client.get(f"/admin/movilidad/acta/{acta.pk}/change/")
    assert r.status_code == 200 and "Pendiente del ToR" in r.content.decode()
    admin_client.post("/admin/movilidad/acta/", {"action": "generar_docx", "_selected_action": [acta.pk]})
    acta.refresh_from_db()
    assert acta.docx
    # corregir pide el motivo en una página intermedia
    acta.firmar(pdf("acta.pdf"), date(2027, 2, 20))
    r = admin_client.post("/admin/movilidad/acta/", {"action": "corregir", "_selected_action": [acta.pk]})
    assert r.status_code == 200 and "Motivo de la corrección" in r.content.decode()
    r = admin_client.post("/admin/movilidad/acta/", {"action": "corregir", "_selected_action": [acta.pk],
                                                     "motivo": "Nota de Databases mal convertida"})
    assert r.status_code == 302 and Acta.objects.filter(corrige=acta).exists()
    assert admin_client.get(f"/admin/movilidad/movilidad/{caso['mov'].pk}/change/").status_code == 200
    assert admin_client.get("/admin/movilidad/escalanotas/").status_code == 200
